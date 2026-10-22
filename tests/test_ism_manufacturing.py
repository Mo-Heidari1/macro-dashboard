"""Tests for the ISM Manufacturing module, run on the committed data files.
The workbook cross-check runs only where the local source workbook exists."""

import os

import numpy as np
import pandas as pd
import pytest

from src.core import paths
from src.core.signals import validate_signal
from src.modules.ism_manufacturing import config as cfg
from src.modules.ism_manufacturing.calculations import (
    classify,
    get_signal,
    heatmap_data,
    load_ism_history,
    load_ism_sectors,
    merge_months,
    peak_note,
    pmi_regime,
    recompute_pmi_from_subindices,
    regime_history,
    sector_rankings,
    summary_table,
)


@pytest.fixture(scope="module")
def committed() -> pd.DataFrame:
    return load_ism_history()


@pytest.fixture(scope="module")
def workbook() -> str:
    path = paths.local_path(cfg.SOURCE_WORKBOOK)
    if not os.path.exists(path):
        pytest.skip("local source workbook not present")
    return path


def _script(name: str):
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        name, os.path.join(paths.ROOT, "scripts", f"{name}.py")
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _pmi_frame(values, start="2024-01-01") -> pd.DataFrame:
    df = pd.DataFrame({"Date": pd.date_range(start, periods=len(values), freq="MS")})
    for col in cfg.SUBINDEX_COLS:
        df[col] = np.nan
    df["PMI"] = values
    return df


# ---------- known values (official release) ----------


@pytest.mark.parametrize(
    "month, expected",
    [
        # Headline PMI as published in ISM's Manufacturing Report On Business;
        # matched against the source workbook's cached values.
        ("2026-09-01", 54.5),
        ("2026-08-01", 54.6),
        ("2026-07-01", 55.6),
        ("2026-06-01", 53.3),
        ("2026-05-01", 54.0),
        ("2026-04-01", 52.7),
        ("2026-01-01", 52.6),
        ("2025-12-01", 47.9),
    ],
)
def test_known_headline_pmi(committed, month, expected):
    assert committed.set_index("Date").loc[month, "PMI"] == pytest.approx(expected)


def test_committed_history_matches_workbook(workbook, committed):
    # every month of every sub-index sheet, plus the Heat Map Summary sheet
    validate = _script("validate_ism")
    series, bad = validate.compare_series(workbook, committed)
    assert series["mismatches"].sum() == 0 and not bad
    assert (series["only in workbook"] == 0).all() and (
        series["only in csv"] == 0
    ).all()
    cells, hm_bad, _ = validate.compare_heatmap(workbook, committed)
    assert cells > 0 and hm_bad == 0


def test_committed_sectors_match_workbook(workbook):
    validate = _script("validate_ism")
    wb_s, _ = validate.read_sectors(workbook)
    wb_s = wb_s[wb_s["Date"] >= pd.Timestamp(cfg.SECTORS_START)]
    m = wb_s.merge(load_ism_sectors(), on=["Category", "Date", "Industry"], how="inner")
    assert len(m) == len(wb_s)  # every workbook row in the window is committed
    assert (m["Direction_x"] == m["Direction_y"]).all()
    assert (m["Score_x"] == m["Score_y"]).all()


def test_committed_history_shape(committed):
    assert list(committed.columns) == ["Date"] + cfg.SUBINDEX_COLS
    assert committed["Date"].min() == pd.Timestamp("1948-01-01")
    assert committed["Date"].is_monotonic_increasing and committed["Date"].is_unique
    assert committed[cfg.SUBINDEX_COLS].iloc[-1].notna().all()  # all 11 series


# ---------- calculations ----------


@pytest.mark.parametrize(
    "value, expected",
    [
        (50.1, "Growth"),
        (50.0, "Neutral"),
        (49.9, "Contraction"),
        (np.nan, "N/A"),
    ],
)
def test_classify_threshold(value, expected):
    assert classify(value) == expected


def test_pmi_weights_sum_to_one():
    assert sum(cfg.PMI_WEIGHTS.values()) == pytest.approx(1.0)


def test_recompute_pmi_known_value():
    df = _pmi_frame([53.0])
    df[["New Orders", "Production", "Employment", "Deliveries", "Inventories"]] = [
        [60, 50, 50, 50, 50]
    ]
    out = recompute_pmi_from_subindices(df).iloc[0]
    assert out["Recomputed PMI"] == pytest.approx(0.30 * 60 + 0.70 * 50)
    assert out["Diff"] == pytest.approx(0.0)


def test_summary_table_mom_yoy():
    df = _pmi_frame(np.arange(48.0, 61.0))  # 13 months, +1 per month
    row = summary_table(df).set_index("Sub-Index").loc["PMI"]
    assert (row["Current"], row["Status"], row["MoM"], row["YoY"]) == (
        60.0,
        "Growth",
        1.0,
        12.0,
    )


def test_heatmap_and_sectors_on_committed_data(committed):
    assert heatmap_data(committed, 12).shape == (12, len(cfg.SUBINDEX_COLS))
    sectors = load_ism_sectors()
    ranked = sector_rankings(sectors, sectors["Date"].max(), "PMI")
    assert not ranked.empty and ranked["Score"].is_monotonic_decreasing


# ---------- get_signal ----------


@pytest.mark.parametrize(
    "values, regime, direction",
    [
        ([52, 52, 52, 54, 55, 56], "Expansion", "improving"),
        ([56, 55, 54, 52, 52, 51], "Slowdown", "deteriorating"),
        ([49, 49, 48, 46, 45.5, 45.5], "Contraction", "deteriorating"),
        ([45, 45, 46, 47, 48, 49], "Recovery", "improving"),
        (
            [52, 52, 52, 52.5, 52.5, 52.5],
            "Expansion",
            "flat",
        ),  # within ±1.0: flat, but rising
        (
            [48, 48, 48, 47.6, 47.6, 47.6],
            "Contraction",
            "flat",
        ),  # within ±1.0: flat, but falling
    ],
)
def test_signal_quadrants(values, regime, direction):
    sig = get_signal(_pmi_frame(values))
    assert (sig["regime"], sig["direction"]) == (regime, direction)


@pytest.mark.parametrize(
    "level, extreme", [(55.1, True), (55.0, False), (45.0, False), (44.9, True)]
)
def test_signal_extreme_flag(level, extreme):
    assert get_signal(_pmi_frame([50, 50, 50, 50, 50, level]))["extreme"] is extreme


def test_pmi_regime_edges():
    assert pmi_regime(50.0, "falling") == "Slowdown"  # exactly 50 counts as above
    assert pmi_regime(49.9, "rising") == "Recovery"
    # first stable month with zero momentum and no prior trend counts as rising
    assert get_signal(_pmi_frame([49.9] * 6))["regime"] == "Recovery"


def test_stable_momentum_keeps_prior_regime():
    # decisive rise (Expansion), then the 3m average stalls and dips by 0.9
    values = [50, 50, 50, 54, 55, 56, 55.5, 55.0, 54.6, 54.2]
    hist = regime_history(_pmi_frame(values))
    assert hist["regime"].iloc[5:8].tolist() == ["Expansion"] * 3
    assert not hist["stable"].iloc[7]
    last = hist.iloc[-1]
    assert last["momentum"] == pytest.approx(-0.9) and bool(last["stable"])
    assert last["regime"] == "Expansion"  # old rule (sign only) gave Slowdown
    sig = get_signal(_pmi_frame(values))
    assert (sig["regime"], sig["direction"]) == ("Expansion", "flat")
    # only 0.8 below its 3-month high (Aug 55.0): under PEAK_NOTE_MIN_DROP
    assert sig["market_read"] == cfg.MARKET_READ_STABLE["Expansion"]


def test_stable_momentum_keeps_trend_when_level_crosses_50():
    hist = regime_history(_pmi_frame([51, 50.6, 50.2, 49.8, 50.4, 50.2, 49.9]))
    assert hist["stable"].iloc[5:].all()
    # first stable month takes the sign of its change (falling); next keeps it
    assert hist["regime"].iloc[5:].tolist() == ["Slowdown", "Contraction"]


def test_decisive_month_uses_regular_market_read():
    sig = get_signal(_pmi_frame([50, 50, 50, 54, 55, 56]))
    assert sig["market_read"] == cfg.MARKET_READ["Expansion"]


def test_signal_skips_month_without_headline():
    df = _pmi_frame([50.0, 51.0, 52.0, np.nan])  # sub-indices typed, PMI not yet
    df.loc[3, "New Orders"] = 55.0
    sig = get_signal(df)
    assert sig["as_of"] == "2024-03-01"
    assert sig["level"] == 52.0 and sig["change"] == 1.0


def test_signal_schema_on_committed_data(committed):
    sig = get_signal(committed)
    assert validate_signal(sig, cfg.DIRECTION_LABELS, cfg.REGIME_LABELS) == []
    assert sig["as_of"] == "2026-09-01" and sig["level"] == 54.5


def test_summary_yoy_uses_calendar_months_across_a_gap():
    df = _pmi_frame([50.0, 51.0, 52.0], start="2025-05-01")
    df.loc[2, "Date"] = pd.Timestamp("2026-06-01")  # 12 months after 2025-06-01
    row = summary_table(df).set_index("Sub-Index").loc["PMI"]
    assert row["YoY"] == pytest.approx(1.0)  # vs June 2025, not the row before
    assert pd.isna(row["MoM"])  # May 2026 missing: no value within the window


def test_summary_reports_each_series_last_month():
    df = _pmi_frame([50.0, 51.0, np.nan])
    df["Prices"] = [60.0, 61.0, 62.0]
    s = summary_table(df).set_index("Sub-Index")
    assert (s.loc["PMI", "As of"], s.loc["PMI", "Current"]) == ("2024-02-01", 51.0)
    assert s.loc["Prices", "As of"] == "2024-03-01"
    assert s.loc["New Orders", "Status"] == "N/A"


def test_sector_rankings_break_ties_by_industry_regardless_of_input_order():
    rows = [
        ("PMI", "2026-06-01", name, "Growth", score)
        for name, score in [("Wood", 3), ("Apparel", 3), ("Paper", 5), ("Chemicals", 3)]
    ]
    df = pd.DataFrame(
        rows, columns=["Category", "Date", "Industry", "Direction", "Score"]
    )
    df["Date"] = pd.to_datetime(df["Date"])
    expected = ["Paper", "Apparel", "Chemicals", "Wood"]
    for seed in range(5):
        shuffled = df.sample(frac=1, random_state=seed)
        assert sector_rankings(shuffled, "2026-06-01")["Industry"].tolist() == expected


def test_zscore_computes_on_full_history(committed):
    assert cfg.ZSCORE_MIN_PERIODS == 120
    z = get_signal(committed)["zscore"]
    assert isinstance(z, float)
    pmi = committed["PMI"].dropna().tail(120)
    assert z == pytest.approx(round((pmi.iloc[-1] - pmi.mean()) / pmi.std(ddof=1), 2))
    assert get_signal(_pmi_frame(list(np.linspace(45, 55, 119))))["zscore"] is None


def test_committed_sectors_only_from_april_2026():
    sectors = load_ism_sectors()
    # no older rows and no comment text: exactly the five ranking columns
    assert list(sectors.columns) == [
        "Category",
        "Date",
        "Industry",
        "Direction",
        "Score",
    ]
    assert sectors["Date"].min() >= pd.Timestamp(cfg.SECTORS_START)
    months = sorted(sectors["Date"].dt.strftime("%Y-%m").unique())
    assert months == ["2026-04", "2026-05", "2026-06", "2026-07", "2026-08", "2026-09"]
    assert sectors.groupby("Date").size().eq(7 * 18).all()


def test_manual_entry_writes_committed_file(tmp_path, monkeypatch, committed):
    update_ism = _script("update_ism")
    assert update_ism.CSV_PATH == paths.dataset_path("ism", "ism_history.csv")
    target = tmp_path / "ism_history.csv"
    committed.to_csv(target, index=False)
    monkeypatch.setattr(update_ism, "CSV_PATH", str(target))
    values = {"Date": "2026-10-01", **{c: 50.0 for c in cfg.SUBINDEX_COLS}}
    update_ism.update(values)
    update_ism.update({**values, "PMI": 51.2})  # same month again: replaces
    out = pd.read_csv(target, parse_dates=["Date"])
    assert len(out) == len(committed) + 1
    assert (
        out.iloc[-1]["PMI"] == 51.2 and out.iloc[-2]["PMI"] == committed.iloc[-1]["PMI"]
    )


def test_merge_months_imports_and_reports():
    current = _pmi_frame([50.0, 51.0, 52.0])  # Jan-Mar 2024
    incoming = _pmi_frame([50.0, 51.5, np.nan, 53.0, 54.0])  # Jan-May 2024
    months = [pd.Timestamp("2024-04-01"), pd.Timestamp("2024-05-01")]
    merged, report = merge_months(
        current, incoming, ["Date"], cfg.SUBINDEX_COLS, months
    )
    assert merged["PMI"].tolist() == [50.0, 51.5, 52.0, 53.0, 54.0]
    kinds = report.groupby("kind").size().to_dict()
    assert kinds == {"imported": 2, "revised": 1, "kept: empty in workbook": 1}
    revised = report[report["kind"] == "revised"].iloc[0]
    assert (revised["Date"], revised["csv"], revised["workbook"]) == (
        pd.Timestamp("2024-02-01"),
        51.0,
        51.5,
    )


def test_merge_months_keeps_rows_missing_from_workbook_and_skips_unimported():
    cols = ["Category", "Date", "Industry", "Direction", "Score"]
    current = pd.DataFrame(
        [("New Orders", "2026-06-01", "Wood", "Growth", 2)], columns=cols
    )
    incoming = pd.DataFrame(
        [
            ("New Orders", "2025-01-01", "Wood", "Growth", 1),
            ("New Orders", "2026-07-01", "Wood", "Contraction", -3),
        ],
        columns=cols,
    )
    for df in (current, incoming):
        df["Date"] = pd.to_datetime(df["Date"])
    merged, report = merge_months(
        current,
        incoming,
        ["Category", "Industry", "Date"],
        ["Direction", "Score"],
        [pd.Timestamp("2026-07-01")],
    )
    assert merged["Date"].dt.strftime("%Y-%m").tolist() == ["2026-06", "2026-07"]
    assert set(report["kind"]) == {
        "imported",
        "kept: not in workbook",
        "skipped: earlier month not in CSV",
    }


def test_extract_sectors_reads_headers_after_a_spacer_column():
    # Layout of the Sep-2026 workbook: Aug 2025 in Q/R, spacer S, Dec 2025 in T/U.
    import datetime
    import importlib.util

    import openpyxl

    spec = importlib.util.spec_from_file_location(
        "extract_ism", os.path.join(paths.ROOT, "scripts", "extract_ism.py")
    )
    extract_ism = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(extract_ism)

    ws = openpyxl.Workbook().active
    ws["B3"], ws["C3"] = "ISM Manufacturing", datetime.datetime(2025, 8, 1)
    # E is a spacer, so Dec 2025 sits in F/G; a fixed C, E, G step skips it.
    ws["F3"] = datetime.datetime(2025, 12, 1)
    ws["B5"], ws["C5"], ws["D5"] = "Wood Products", "Growth", 3
    ws["F5"], ws["G5"] = "Contraction", -2
    out = extract_ism.extract_sectors(ws)
    assert out["Date"].dt.strftime("%Y-%m").tolist() == ["2025-08", "2025-12"]
    assert out["Score"].tolist() == [3, -2]


# ---------- peak note in market_read ----------


def test_market_read_notes_drop_from_three_month_peak(committed):
    # Jul-Sep 2026 published PMI: 55.6, 54.6, 54.5; momentum is rising.
    sig = get_signal(committed)
    assert sig["market_read"].startswith(
        "3-month average rising, but off the July peak (55.6). "
    )
    assert sig["market_read"].endswith(cfg.MARKET_READ["Expansion"])
    # text only: regime, direction and thresholds are untouched
    assert (sig["regime"], sig["direction"]) == ("Expansion", "improving")
    assert (cfg.DIRECTION_BAND, cfg.DIRECTION_WINDOW, cfg.REGIME_LEVEL) == (
        1.0,
        3,
        50.0,
    )
    assert (cfg.EXTREME_HIGH, cfg.EXTREME_LOW) == (55.0, 45.0)


def test_no_peak_note_at_a_new_high():
    sig = get_signal(_pmi_frame([50, 50, 50, 54, 55, 56]))
    assert sig["market_read"] == cfg.MARKET_READ["Expansion"]


def test_peak_note_wording_when_falling_and_when_stable():
    falling = get_signal(_pmi_frame([56, 55, 54, 52, 52.5, 51]))
    assert falling["regime"] == "Slowdown"
    assert falling["market_read"].startswith(
        "3-month average falling, and off the May peak (52.5). "
    )
    # momentum 0.0 (stable, trend carried as rising); 1.1 below the Aug peak
    stable = get_signal(_pmi_frame([50, 50, 50, 54, 55, 56, 55.5, 56.0, 55.6, 54.9]))
    assert (stable["regime"], stable["direction"]) == ("Expansion", "flat")
    assert stable["market_read"] == (
        "Off the August peak (56.0). " + cfg.MARKET_READ_STABLE["Expansion"]
    )


def test_peak_note_uses_latest_tied_peak_and_names_the_year_across_january():
    tied = peak_note(
        pd.Series(pd.date_range("2026-07-01", periods=3, freq="MS")),
        pd.Series([55.0, 55.0, 54.0]),
        "rising",
        False,
    )
    assert tied == "3-month average rising, but off the August peak (55.0)."
    across = peak_note(
        pd.Series(pd.date_range("2026-11-01", periods=3, freq="MS")),
        pd.Series([53.0, 52.0, 51.5]),
        "falling",
        False,
    )
    assert across == "3-month average falling, and off the November 2026 peak (53.0)."


@pytest.mark.parametrize(
    "latest, noted",
    [
        (54.7, False),  # 0.9 below the 55.6 peak: just under the cut, no note
        (54.6, True),  # exactly 1.0 below: at the cut, noted
        (54.5, True),  # 1.1 below: just over the cut, noted
    ],
)
def test_peak_note_needs_minimum_drop(latest, noted):
    assert cfg.PEAK_NOTE_MIN_DROP == 1.0
    dates = pd.Series(pd.date_range("2026-07-01", periods=3, freq="MS"))
    note = peak_note(dates, pd.Series([55.6, 55.4, latest]), "rising", False)
    assert (note is not None) is noted
    if noted:
        assert note == "3-month average rising, but off the July peak (55.6)."


def test_small_pullback_keeps_plain_market_read():
    # momentum rising decisively; latest 0.3 below the 3-month high
    sig = get_signal(_pmi_frame([50, 50, 50, 54, 55.3, 55.0]))
    assert sig["regime"] == "Expansion"
    assert sig["market_read"] == cfg.MARKET_READ["Expansion"]
